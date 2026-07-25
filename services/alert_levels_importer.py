# services/alert_levels_importer.py
"""
Lee una planilla (.xlsx / .xlsm / .csv) con niveles de alarma y devuelve
filas normalizadas listas para upsert_alert().

Formato esperado (la primera fila con encabezados reconocibles manda):

    symbol | target | stop_loss | enabled
    EWW    | 76     | 74        | si

Los encabezados admiten variantes en castellano e inglés y no distinguen
mayúsculas ni acentos. Las columnas que sobran se ignoran.

API pública:
    parse_levels(file_bytes, filename) -> list[dict]
        [{'symbol': 'EWW', 'target': 76.0, 'stop_loss': 74.0,
          'enabled': True, 'row': 2, 'error': None}, ...]
"""
from __future__ import annotations

import csv
import io
import re
import unicodedata

import openpyxl


# ── Alias de encabezados ───────────────────────────────────────────────
_ALIASES = {
    'symbol':    {'symbol', 'symbols', 'simbolo', 'ticker', 'tickers',
                  'activo', 'activos', 'asset', 'especie', 'papel'},
    'target':    {'target', 'targetprice', 'target price', 'takeprofit',
                  'take profit', 'tp', 'objetivo', 'precio objetivo',
                  'precio target', 'toma de ganancia'},
    'stop_loss': {'stoploss', 'stop loss', 'stop', 'sl', 'stopprice',
                  'stop price', 'precio stop', 'perdida maxima'},
    'enabled':   {'enabled', 'activa', 'alarma', 'habilitada', 'on', 'estado'},
}

_TRUE = {'1', 'true', 'si', 'sí', 'yes', 'y', 'x', 'on', 'activa', 'activo',
         'verdadero', 'ok'}
_FALSE = {'0', 'false', 'no', 'n', 'off', 'inactiva', 'inactivo', 'falso'}


def _norm(v) -> str:
    """minúsculas, sin acentos, sin dobles espacios."""
    s = '' if v is None else str(v)
    s = unicodedata.normalize('NFKD', s)
    s = ''.join(c for c in s if not unicodedata.combining(c))
    s = re.sub(r'[\s_\-\.]+', ' ', s).strip().lower()
    return s


def _match_header(cell) -> str | None:
    n = _norm(cell)
    if not n:
        return None
    compact = n.replace(' ', '')
    for field, names in _ALIASES.items():
        for name in names:
            if n == name or compact == name.replace(' ', ''):
                return field
    return None


def _to_float(v):
    """Acepta 76, '76', '76,5', '$76.50', '  76 '. Vacío -> None."""
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip()
    if not s:
        return None
    s = s.replace('$', '').replace('%', '').replace(' ', '')
    # 1.234,56 -> 1234.56  |  1,234.56 -> 1234.56  |  76,5 -> 76.5
    if ',' in s and '.' in s:
        s = s.replace('.', '').replace(',', '.') if s.rfind(',') > s.rfind('.') \
            else s.replace(',', '')
    elif ',' in s:
        s = s.replace(',', '.')
    try:
        return float(s)
    except ValueError:
        raise ValueError(f"'{v}' no es un número válido")


def _to_bool(v, default=True):
    if v is None or str(v).strip() == '':
        return default
    n = _norm(v)
    if n in _TRUE:
        return True
    if n in _FALSE:
        return False
    return default


# ── Lectura de la grilla ───────────────────────────────────────────────
def _grid_from_xlsx(file_bytes: bytes) -> list[list]:
    wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True, read_only=True)
    ws = wb[wb.sheetnames[0]]
    grid = [list(row) for row in ws.iter_rows(values_only=True)]
    wb.close()
    return grid


def _grid_from_csv(file_bytes: bytes) -> list[list]:
    text = file_bytes.decode('utf-8-sig', errors='ignore')
    sample = text[:2048]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=',;\t')
    except csv.Error:
        dialect = csv.excel
    return [row for row in csv.reader(io.StringIO(text), dialect)]


def _find_header(grid: list[list]) -> tuple[int, dict]:
    """Devuelve (índice de la fila de encabezados, {campo: índice de columna})."""
    for i, row in enumerate(grid[:15]):
        mapping = {}
        for j, cell in enumerate(row):
            field = _match_header(cell)
            if field and field not in mapping:
                mapping[field] = j
        if 'symbol' in mapping and ('target' in mapping or 'stop_loss' in mapping):
            return i, mapping
    raise ValueError(
        "No se encontró la fila de encabezados. La planilla necesita una "
        "columna 'symbol' y al menos una de 'target' o 'stop_loss'.")


# ── API pública ────────────────────────────────────────────────────────
def parse_levels(file_bytes: bytes, filename: str = '') -> list[dict]:
    name = (filename or '').lower()
    grid = _grid_from_csv(file_bytes) if name.endswith('.csv') \
        else _grid_from_xlsx(file_bytes)

    if not grid:
        raise ValueError("La planilla está vacía")

    header_idx, cols = _find_header(grid)
    out, seen = [], set()

    started = False
    for offset, row in enumerate(grid[header_idx + 1:], start=header_idx + 2):
        def cell(field):
            j = cols.get(field)
            return row[j] if j is not None and j < len(row) else None

        # Una fila totalmente vacía después de los datos corta la lectura:
        # todo lo que venga abajo (notas, leyendas) se ignora.
        if all(str(c or '').strip() == '' for c in row):
            if started:
                break
            continue

        symbol = str(cell('symbol') or '').strip().upper()
        if not symbol:
            continue
        started = True
        if len(symbol) > 20:
            out.append({'symbol': symbol[:20], 'row': offset,
                        'error': 'símbolo demasiado largo'})
            continue
        if symbol in seen:
            out.append({'symbol': symbol, 'row': offset,
                        'error': 'símbolo repetido en la planilla'})
            continue
        seen.add(symbol)

        entry = {'symbol': symbol, 'row': offset, 'error': None,
                 'target': None, 'stop_loss': None, 'enabled': True}
        try:
            entry['target'] = _to_float(cell('target'))
            entry['stop_loss'] = _to_float(cell('stop_loss'))
        except ValueError as e:
            entry['error'] = str(e)
            out.append(entry)
            continue

        entry['enabled'] = _to_bool(cell('enabled'), default=True)

        if entry['target'] is None and entry['stop_loss'] is None:
            entry['error'] = 'sin target ni stop loss'
        elif (entry['target'] is not None and entry['stop_loss'] is not None
              and entry['stop_loss'] >= entry['target']):
            entry['error'] = 'el stop loss tiene que estar por debajo del target'

        out.append(entry)

    if not out:
        raise ValueError("No se encontró ninguna fila con datos debajo de los encabezados")
    return out
